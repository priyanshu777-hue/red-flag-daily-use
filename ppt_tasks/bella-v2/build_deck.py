"""Bella by Red Flag Homes — v2 deck (6 slides). Quiet luxury per the v2 design system: black #0B0A0C, cream #F4EFE8,
champagne gold #C9A66B for lines/details, red #C8102E only for buttons and key numbers; Cormorant Garamond + DM Sans.
Bella is shown as an original AI-avatar orb (make_art.py). Build Mode, pptx-designer public API.
Run: python make_art.py && python build_deck.py
"""
from pathlib import Path

from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from pptx_designer import Presentation
from pptx_designer.tools.layout import add_slide, clean_save
from pptx_designer.tools.shapes import oval, rect, rrect
from pptx_designer.tools.text import text

import motion

HERE = Path(__file__).resolve().parent
ART = HERE / "art"
OUT = HERE / "output"

SERIF, SANS = "Cormorant Garamond", "DM Sans"
W, H, M = 13.333, 7.5, 0.85
BLACK, CREAM, GOLD, RED, WHITE = "#0B0A0C", "#F4EFE8", "#C9A66B", "#C8102E", "#FFFFFF"
MUTE_D, MUTE_L, CARD_D, CARD_L, LINE_D, LINE_L = "#8F8890", "#7C7470", "#141215", "#FBF8F3", "#2E2A2F", "#DDD3C4"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.bg = BLACK

    @property
    def C(self):
        return {"background": self.bg, "text_body": CREAM, "text_dark": CREAM, "text_muted": MUTE_D,
                "font_body": SANS, "font_heading": SERIF}

    def T(self, s, x, y, w, h, txt, size=14, font=SANS, color=CREAM, bold=False, italic=False, align="left",
          spacing=None, leading=None, caps=False, strike=False):
        box = text(s, x, y, w, h, txt.upper() if caps else txt, font_size=size, color=color, bold=bold,
                   align=align, font_name=font, C=self.C)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for p in tf.paragraphs:
            p.alignment = tf.paragraphs[0].alignment
            if leading:
                p.line_spacing = leading
            for r in p.runs:
                r.font.italic = italic
                if spacing is not None:
                    r.font._element.set("spc", str(spacing))
                if strike:
                    r.font._element.set("strike", "sngStrike")
        return box

    def serif(self, s, x, y, w, h, txt, size=44, color=CREAM, italic=False, bold=False, align="left", leading=0.95):
        return self.T(s, x, y, w, h, txt, size=size, font=SERIF, color=color, bold=bold, italic=italic, align=align,
                      leading=leading)

    def label(self, s, x, y, w, txt, color=GOLD, align="left", size=8.5):
        return self.T(s, x, y, w, 0.25, txt, size=size, color=color, bold=True, caps=True, spacing=350, align=align)

    def hair(self, s, x, y, w, color=GOLD):
        return rect(s, x, y, w, 0.01, color, C=self.C)

    def glow(self, shp, color="C9A66B", rad=18, alpha=35):
        sp = shp._element.spPr
        eff = etree.SubElement(sp, qn("a:effectLst"))
        g = etree.SubElement(eff, qn("a:glow"), rad=str(rad * 12700))
        c = etree.SubElement(g, qn("a:srgbClr"), val=color)
        etree.SubElement(c, qn("a:alpha"), val=str(alpha * 1000))
        return shp

    def card(self, s, x, y, w, h, fill, line, radius=0.08, lw=0.75):
        b = rrect(s, x, y, w, h, fill, line=line, C=self.C)
        b.adjustments[0] = radius
        b.line.width = Pt(lw)
        return b

    def button(self, s, x, y, w, h, label, size=14):
        b = self.card(s, x, y, w, h, RED, RED, radius=0.5)
        self.glow(b, "C8102E", 14, 30)
        self.T(s, x, y + (h - size / 72 * 1.3) / 2, w, h, label, size=size, color=WHITE, bold=True, align="center")
        return b

    def pic(self, s, name, x, y, w=None, h=None, tag="static-art"):
        p = s.shapes.add_picture(str(ART / name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
        p.name = tag
        return p

    def slide(self, dark=True, notes=None):
        self.bg = BLACK if dark else CREAM
        s = add_slide(self.prs)
        self.pic(s, "grain_black.jpg" if dark else "grain_cream.jpg", 0, 0, w=W, h=H, tag="static-bg")
        n = len(self.prs.slides)
        for shp in (self.label(s, M, H - 0.55, 5, "Bella  ·  Red Flag Homes", color=MUTE_D if dark else MUTE_L, size=7),
                    self.T(s, W - M - 1, H - 0.55, 1, 0.25, f"0{n} / 06", size=7.5, color=GOLD, bold=True, align="right",
                           spacing=200)):
            shp.name = "chrome"
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self):
        for i, sl in enumerate(self.prs.slides, start=1):
            motion.transition(sl, "fade", through_black=i in (1, 4, 5))
            motion.choreograph(sl, step_ms=140, budget_ms=2400)
        OUT.mkdir(exist_ok=True)
        path = OUT / "Bella_v2.pptx"
        clean_save(self.prs, str(path))
        print(path)


def build():
    d = Deck()
    T, serif, label, hair = d.T, d.serif, d.label, d.hair
    CLOSE = ("Speaker line for the close: “Your guests don’t wait, and neither should your bnb. Try Bella free for 48 hours. "
             "If you don’t love it, cancel anytime.”")

    # 1 — Cover --------------------------------------------------------------------------------------
    s = d.slide(notes="Open on the hook: your bnb never sleeps.")
    d.pic(s, "logo.png", M - 0.08, 0.45, w=0.85)
    label(s, W - M - 4, 0.62, 4, "Powered by GPT-5.6 Sol", align="right")
    d.pic(s, "orb.png", 7.35, 0.55, w=5.6)
    serif(s, M - 0.08, 1.75, 7, 2.6, "Bella", size=190, italic=True, leading=0.85)
    hair(s, M, 4.38, 3.6)
    label(s, M, 4.52, 6, "by Red Flag Homes", color=GOLD)
    serif(s, M, 4.95, 6.3, 1.0, "Your bnb never sleeps.\nNow neither does Bella.", size=28, leading=1.05)
    T(s, M, 6.0, 5.4, 0.5, "The AI co-host that replies, prices, checks in and protects your calendar. 24/7. In 50+ "
      "languages.", size=10.5, color=MUTE_D, leading=1.4)
    d.button(s, 7.85, 5.85, 4.6, 0.62, "Start your 48-hour free trial  →", size=13.5)
    T(s, 7.85, 6.6, 4.6, 0.25, "Cancel anytime  ·  No commission  ·  Festive offer ₹2,999 / year", size=8.5,
      color=GOLD, align="center")

    # 2 — See Bella in action (cream) --------------------------------------------------------------------
    s = d.slide(dark=False, notes="Let the chat speak. 2 AM, Hinglish, three seconds.")
    label(s, M, 1.1, 6, "See Bella in action")
    serif(s, M, 1.5, 6.4, 2.2, "A guest texts at 2 AM.\nBella answers in 3 seconds.", size=44, color=BLACK, leading=1.0)
    hair(s, M, 3.85, 1.2)
    T(s, M, 4.1, 5.4, 0.9, "Same question in any language. Bella replies in the guest’s own language.", size=13,
      color="#4A4346", leading=1.45)
    for i, (big, small) in enumerate((("3s", "average reply"), ("50+", "languages"), ("24/7", "even at 2 AM"))):
        x = M + i * 2.15
        hair(s, x, 5.05, 1.8, LINE_L)
        serif(s, x, 5.2, 1.9, 0.8, big, size=44, color=RED if i == 0 else BLACK)
        T(s, x, 6.0, 1.9, 0.3, small, size=10, color=MUTE_L)
    px, py, pw, ph = 8.15, 0.75, 3.9, 6.1
    phone = d.card(s, px, py, pw, ph, BLACK, BLACK, radius=0.12)
    d.card(s, px + 0.14, py + 0.14, pw - 0.28, ph - 0.28, CARD_L, CARD_L, radius=0.1)
    d.card(s, px + pw / 2 - 0.55, py + 0.27, 1.1, 0.22, BLACK, BLACK, radius=0.5)
    d.pic(s, "orb.png", px + 0.3, py + 0.62, w=0.75)
    T(s, px + 1.08, py + 0.77, 2.5, 0.3, "Bella", size=13, color=BLACK, bold=True)
    T(s, px + 1.08, py + 1.03, 2.6, 0.25, "AI co-host  ·  replies in 3s", size=8.5, color=GOLD, bold=True)
    hair(s, px + 0.14, py + 1.42, pw - 0.28, LINE_L)
    T(s, px, py + 1.6, pw, 0.25, "2:04 AM", size=8.5, color=MUTE_L, bold=True, align="center")
    d.card(s, px + 0.35, py + 2.0, 2.75, 0.62, "#E7E1D8", "#E7E1D8", radius=0.35)
    T(s, px + 0.55, py + 2.17, 2.5, 0.3, "India Gate kitni door hai?", size=11.5, color=BLACK)
    T(s, px + 0.35, py + 2.68, 2.6, 0.2, "Guest", size=7.5, color=MUTE_L)
    bb = d.card(s, px + pw - 2.6 - 0.35, py + 3.05, 2.6, 0.62, BLACK, GOLD, radius=0.35, lw=1.25)
    d.glow(bb, "C9A66B", 8, 25)
    T(s, px + pw - 2.6 - 0.15, py + 3.22, 2.3, 0.3, "Sirf 10 min door.", size=11.5, color=CREAM)
    T(s, px + pw - 2.95, py + 3.73, 2.6, 0.2, "Bella  ·  3 seconds", size=7.5, color=GOLD, align="right", bold=True)
    for k in range(3):
        oval(s, px + 0.5 + k * 0.2, py + 4.35, 0.1, 0.1, "#CFC6B8", C=d.C)
    T(s, px + 0.35, py + ph - 0.65, pw - 0.7, 0.25, "Illustrative conversation", size=7.5, color=MUTE_L, align="center")

    # 3 — Everything handled (black) ------------------------------------------------------------------
    s = d.slide(notes="Six things, one co-host. Close on: live in 3 clicks.")
    label(s, M, 0.85, 6, "Everything handled")
    serif(s, M, 1.15, 11, 0.9, "Everything your listing needs, handled.", size=40)
    feats = [("i_reply", "Instant guest replies", "Answers every guest in seconds, day or night, in 50+ languages including Hinglish."),
             ("i_price", "Smart pricing", "Adjusts your nightly rates for weekends, festivals and demand."),
             ("i_key", "Automatic check-in", "Sends every guest their check-in details on time, every time."),
             ("i_cal", "Calendar blocking", "Blocks your dates for you, so no double bookings."),
             ("i_audit", "A to Z listing audit", "Reviews photos, title, description and settings, then fixes the gaps."),
             ("i_cohost", "Co-Host chat", "Want anything else? Just type it to Bella in the Co-Host section and it’s done.")]
    cw, ch, g = (W - 2 * M - 0.5) / 3, 1.95, 0.25
    for i, (ic, t, b) in enumerate(feats):
        x = M + (i % 3) * (cw + g)
        y = 2.25 + (i // 3) * (ch + g)
        d.card(s, x, y, cw, ch, CARD_D, LINE_D)
        d.pic(s, f"{ic}.png", x + 0.32, y + 0.3, w=0.48)
        serif(s, x + 0.95, y + 0.33, cw - 1.1, 0.5, t, size=21)
        T(s, x + 0.32, y + 0.98, cw - 0.64, 0.85, b, size=10.5, color=MUTE_D, leading=1.4)
    hair(s, M, 6.65, 0.6)
    T(s, M + 0.8, 6.55, 7, 0.3, "Zero setup hassle.  Live in 3 clicks.", size=11, color=GOLD, bold=True)

    # 4 — The brain (cream) --------------------------------------------------------------------------
    s = d.slide(dark=False, notes="Say it slowly. One statement, then the three benefits.")
    d.pic(s, "glow.png", 7.0, 0.6, w=5.0)
    label(s, 0, 1.3, W, "The brain behind Bella", align="center")
    serif(s, 0, 2.15, W, 1.4, "Powered by GPT-5.6 Sol.", size=72, color=BLACK, align="center")
    hair(s, W / 2 - 0.6, 3.75, 1.2)
    T(s, 2.2, 4.0, W - 4.4, 0.5, "The most powerful AI model in the world right now, working only for your bnb.", size=15,
      color="#4A4346", align="center")
    serif(s, 0, 4.85, W, 0.6, "Smarter replies.  Better pricing decisions.  Fewer mistakes.", size=24, italic=True,
          color=BLACK, align="center")

    # 5 — Pricing (black) ---------------------------------------------------------------------------
    s = d.slide(notes="Struck price first, then the red card. Land on ≈ ₹125 a month per listing.")
    label(s, M, 0.85, 6, "Pricing")
    serif(s, M, 1.15, 11, 0.9, "No commission. Just one yearly fee.", size=40)
    lx, ly, lw, lh = M, 2.4, 4.6, 3.3
    d.card(s, lx, ly, lw, lh, CARD_D, LINE_D)
    label(s, lx + 0.4, ly + 0.4, 3, "Regular price", color=MUTE_D)
    T(s, lx + 0.4, ly + 0.85, 4, 0.9, "₹4,999", size=50, color=MUTE_D, strike=True)
    T(s, lx + 0.4, ly + 1.75, 3, 0.3, "per year", size=12, color=MUTE_D)
    hair(s, lx + 0.4, ly + 2.3, lw - 0.8, LINE_D)
    T(s, lx + 0.4, ly + 2.45, lw - 0.8, 0.6, "Up to 2 listings\n48-hour free trial", size=11, color=MUTE_D, leading=1.4)
    rx, ry, rw, rh = 5.85, 2.1, 4.9, 3.9
    rc = d.card(s, rx, ry, rw, rh, RED, GOLD, lw=1.25)
    d.glow(rc, "C8102E", 22, 28)
    label(s, rx + 0.45, ry + 0.45, 3, "Festive offer", color="#F3DFAE")
    T(s, rx + 0.4, ry + 0.85, 4.3, 1.2, "₹2,999", size=78, color=WHITE, bold=True)
    T(s, rx + 0.45, ry + 2.1, 3, 0.3, "per year", size=14, color=WHITE)
    hair(s, rx + 0.45, ry + 2.65, rw - 0.9, "#E7A3AE")
    T(s, rx + 0.45, ry + 2.85, rw - 0.9, 0.6, "Up to 2 listings\n48-hour free trial", size=12.5, color=WHITE, bold=True,
      leading=1.4)
    d.pic(s, "orb.png", 10.85, 2.0, w=2.0)
    T(s, 10.95, 4.1, 1.8, 0.6, "Exclusive to Red Flag’s first clients.", size=9.5, color=GOLD, align="center", leading=1.3)
    T(s, M, 6.25, 11, 0.35, "Cancel anytime", size=12, color=GOLD, bold=True)
    T(s, M + 1.7, 6.26, 10, 0.35, "That’s about ₹125 a month per listing with two listings. Start with a 48-hour free trial.",
      size=11.5, color=CREAM)

    # 6 — How to start (cream) -------------------------------------------------------------------------
    s = d.slide(dark=False, notes=CLOSE)
    label(s, M, 0.85, 6, "How to start")
    serif(s, M, 1.15, 11, 0.9, "Three steps. Then Bella takes over.", size=40, color=BLACK)
    steps = [("01", "Start your free trial", "48 hours free, up to 2 listings. Cancel anytime.", None),
             ("02", "Add Bella as co-host", "Invite her with full access:", "bella@redflaghomes.in"),
             ("03", "Save. Done.", "Bella now manages everything.", None)]
    sw = (W - 2 * M - 0.5) / 3
    for i, (n, t, b, chip) in enumerate(steps):
        x = M + i * (sw + 0.25)
        y = 2.35
        d.card(s, x, y, sw, 2.6, CARD_L, LINE_L)
        serif(s, x + 0.35, y + 0.25, 1.5, 0.9, n, size=54, color=GOLD)
        hair(s, x + 0.35, y + 1.2, sw - 0.7, LINE_L)
        serif(s, x + 0.35, y + 1.35, sw - 0.7, 0.5, t, size=24, color=BLACK)
        T(s, x + 0.35, y + 1.88, sw - 0.7, 0.5, b, size=10.5, color="#4A4346", leading=1.35)
        if chip:
            d.card(s, x + 0.35, y + 2.2, 2.55, 0.3, BLACK, GOLD, radius=0.5, lw=0.75)
            T(s, x + 0.35, y + 2.25, 2.55, 0.25, chip, size=9, color=CREAM, bold=True, align="center")
    rect(s, 0, 5.4, W, 2.1, BLACK, C=d.C)
    d.pic(s, "orb.png", M - 0.15, 5.5, w=1.6)
    serif(s, M + 1.6, 5.75, 6.5, 0.6, "Your guests don’t wait.", size=28, italic=True)
    T(s, M + 1.6, 6.3, 6.5, 0.3, "Neither should your bnb. Try Bella free for 48 hours.", size=11, color=MUTE_D)
    d.button(s, W - M - 3.2, 5.85, 3.2, 0.72, "Get Bella  →", size=17)
    for shp in [x for x in s.shapes]:                                  # chrome over the black strip
        if shp.name == "chrome":
            s.shapes._spTree.remove(shp._element)
            s.shapes._spTree.append(shp._element)

    d.save()


if __name__ == "__main__":
    build()
