"""Motion layer for the Visibility Booster deck.

Slide transitions go through pptx_designer.effects.animation.add_slide_transition, then
``_fix_order`` moves <p:transition> after <p:clrMapOvr>. The library puts it straight after
<p:cSld>, which breaks the p:sld schema order, and PowerPoint then offers to repair the file.

Entrance builds are written here as standard PowerPoint timing XML. The library's
add_entrance_animation only toggles visibility (no fade/float behaviour) and starts each build on a
non-click condition, so it cannot produce a real cascade. Morph is written as the
mc:AlternateContent block PowerPoint itself saves (p159:morph, with a fade fallback for older
viewers and LibreOffice).
"""
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Emu

from pptx_designer.effects.animation import add_slide_transition

P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"

STATIC = "static"   # name prefix: never animated (photos, masks, veils, backgrounds)
CHROME = "chrome"   # name prefix: running header/footer/mark — present from the first frame


# --------------------------------------------------------------------------- transitions
def _fix_order(sld):
    """Re-seat transition/timing in schema order: cSld, clrMapOvr, transition, timing, extLst."""
    order = [qn("p:cSld"), qn("p:clrMapOvr"), qn("p:transition"), f"{{{MC}}}AlternateContent",
             qn("p:timing"), qn("p:extLst")]
    kids = list(sld)
    for k in kids:
        sld.remove(k)
    for k in sorted(kids, key=lambda e: order.index(e.tag) if e.tag in order else 99):
        sld.append(k)


def transition(slide, kind="fade", through_black=False, speed="slow"):
    sld = slide._element
    if kind == "morph":
        for old in sld.findall(f"{{{MC}}}AlternateContent") + sld.findall(qn("p:transition")):
            sld.remove(old)
        xml = (
            f'<mc:AlternateContent xmlns:mc="{MC}">'
            f'<mc:Choice xmlns:p159="{P159}" Requires="p159">'
            f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            f'xmlns:p14="{P14}" spd="slow" p14:dur="1600"><p159:morph option="byObject"/></p:transition>'
            f'</mc:Choice><mc:Fallback>'
            f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow">'
            f'<p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
        )
        sld.append(etree.fromstring(xml))
    else:
        add_slide_transition(slide, transition_type=kind, speed=speed)
        if through_black:
            sld.find(qn("p:transition"))[0].set("thruBlk", "1")
    _fix_order(sld)


# --------------------------------------------------------------------------- entrance builds
class _Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return str(self.n)


def _E(tag, **attrs):
    el = etree.Element(qn(tag))
    for k, v in attrs.items():
        el.set(k, v)
    return el


def _sub(parent, tag, **attrs):
    el = _E(tag, **attrs)
    parent.append(el)
    return el


def _cbhvr(parent, ids, spid, dur, attr=None, fill=None, extra=None):
    b = _sub(parent, "p:cBhvr")
    ctn = _sub(b, "p:cTn", id=ids(), dur=str(dur), **({"fill": fill} if fill else {}))
    for k, v in (extra or {}).items():
        ctn.set(k, v)
    if dur == 1:
        st = _sub(ctn, "p:stCondLst")
        _sub(st, "p:cond", delay="0")
    tgt = _sub(b, "p:tgtEl")
    _sub(tgt, "p:spTgt", spid=str(spid))
    if attr:
        lst = _sub(b, "p:attrNameLst")
        _sub(lst, "p:attrName").text = attr
    return b


def _effect(parent, ids, spid, kind, delay, dur, first):
    # presetID: 10 = Fade, 42 = Float In (Ascend), 22 = Wipe (subtype 8 = from left)
    preset, subtype = {"fade": ("10", "0"), "float": ("42", "0"), "wipe": ("22", "8")}[kind]
    par = _sub(parent, "p:par")
    ctn = _sub(par, "p:cTn", id=ids(), presetID=preset, presetClass="entr", presetSubtype=subtype,
               fill="hold", grpId="0", nodeType="afterEffect" if first else "withEffect")
    st = _sub(ctn, "p:stCondLst")
    _sub(st, "p:cond", delay=str(delay))
    kids = _sub(ctn, "p:childTnLst")

    s = _sub(kids, "p:set")
    _cbhvr(s, ids, spid, 1, "style.visibility", fill="hold")
    to = _sub(s, "p:to")
    _sub(to, "p:strVal", val="visible")

    flt = "wipe(left)" if kind == "wipe" else "fade"
    ae = _sub(kids, "p:animEffect", transition="in", filter=flt)
    _cbhvr(ae, ids, spid, dur)

    if kind == "float":
        for attr, frm, to_ in (("ppt_x", "#ppt_x", "#ppt_x"), ("ppt_y", "#ppt_y+.035", "#ppt_y")):
            an = _sub(kids, "p:anim", calcmode="lin", valueType="num")
            _cbhvr(an, ids, spid, dur, attr, extra={"decel": "100000"})
            tl = _sub(an, "p:tavLst")
            for tm, val in (("0", frm), ("100000", to_)):
                tav = _sub(tl, "p:tav", tm=tm)
                v = _sub(tav, "p:val")
                _sub(v, "p:strVal", val=val)


def _kind_for(shape):
    """Pick an effect from what the shape is: hairlines draw in, big type rises, the rest fades."""
    w, h = Emu(shape.width).inches, Emu(shape.height).inches
    if not shape.has_text_frame or not shape.text_frame.text.strip():
        return "wipe" if min(w, h) < 0.05 and w > h else "fade"
    size = max((r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size),
               default=12)
    return "float" if size >= 17 else "fade"


def choreograph(slide, step_ms=110, budget_ms=1900):
    """Auto-play cascade on slide entry: content reveals in reading order (rows, then columns)."""
    shapes = [s for s in slide.shapes
              if not s.name.startswith((STATIC, CHROME)) and s.shape_type != 13]  # 13 = picture
    if not shapes:
        return
    row = lambda s: round(Emu(s.top).inches / 0.35)  # noqa: E731 — bucket near-equal tops together
    shapes.sort(key=lambda s: (row(s), Emu(s.left).inches))
    step = min(step_ms, budget_ms // max(1, len(shapes)))

    sld = slide._element
    for old in sld.findall(qn("p:timing")):
        sld.remove(old)
    ids = _Ids()
    timing = _E("p:timing")
    tn = _sub(timing, "p:tnLst")
    root = _sub(_sub(tn, "p:par"), "p:cTn", id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    seq = _sub(_sub(root, "p:childTnLst"), "p:seq", concurrent="1", nextAc="seek")
    main = _sub(seq, "p:cTn", id="2", dur="indefinite", nodeType="mainSeq")
    main_kids = _sub(main, "p:childTnLst")
    for evt, lst in (("onPrev", "p:prevCondLst"), ("onNext", "p:nextCondLst")):
        c = _sub(_sub(seq, lst), "p:cond", evt=evt, delay="0")
        _sub(_sub(c, "p:tgtEl"), "p:sldTgt")

    # One build group that starts automatically when the slide begins.
    grp = _sub(_sub(main_kids, "p:par"), "p:cTn", id=ids(), fill="hold")
    st = _sub(grp, "p:stCondLst")
    _sub(st, "p:cond", delay="indefinite")
    on_begin = _sub(st, "p:cond", evt="onBegin", delay="0")
    _sub(on_begin, "p:tn", val="2")
    inner = _sub(_sub(_sub(grp, "p:childTnLst"), "p:par"), "p:cTn", id=ids(), fill="hold")
    _sub(_sub(inner, "p:stCondLst"), "p:cond", delay="0")
    inner_kids = _sub(inner, "p:childTnLst")

    bld = _E("p:bldLst")
    for i, shp in enumerate(shapes):
        kind = _kind_for(shp)
        dur = {"float": 900, "wipe": 700, "fade": 700}[kind]
        _effect(inner_kids, ids, shp.shape_id, kind, 250 + i * step, dur, first=(i == 0))
        if shp.has_text_frame:
            _sub(bld, "p:bldP", spid=str(shp.shape_id), grpId="0", animBg="1")
    if len(bld):
        timing.append(bld)
    sld.append(timing)
    _fix_order(sld)
