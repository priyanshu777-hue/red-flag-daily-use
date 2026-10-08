"""Build the Bella by Red Flag Homes deck — 4 slides, per the brief's look:
black #141214, cream #F4EFE8, Red Flag red #C8102E; Cormorant Garamond headings, DM Sans body.
Pinterest-editorial moves: giant italic name, arch portrait, sticker badges, sparkles, a 3 AM chat mockup,
price-tag cards. Build Mode, pptx-designer public API. Run: python make_art.py && python build_deck.py
"""
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from pptx_designer import Presentation
from pptx_designer.tools.layout import add_slide, clean_save
from pptx_designer.tools.shapes import oval, rect, rrect, shape
from pptx_designer.tools.text import text

import motion

HERE = Path(__file__).resolve().parent
ART = HERE / "art"
OUT = HERE / "output"

SERIF, SANS = "Cormorant Garamond", "DM Sans"
W, H, M = 13.333, 7.5, 0.75
BLACK, CREAM, RED = "#141214", "#F4EFE8", "#C8102E"
INK, MUTE, MUTE_D, CARD_D, HAIR_D, HAIR_L = "#141214", "#7A7276", "#9C9498", "#211E21", "#3A3438", "#E2DAD0"
GOLD, BLUSH, WHITE = "#E9B872", "#F2D3CF", "#FFFFFF"


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.bg = BLACK

    @property
    def C(self):
        return {"background": self.bg, "text_body": CREAM, "text_dark": CREAM, "text_muted": MUTE_D,
                "font_body": SANS, "font_heading": SERIF}

    def T(self, s, x, y, w, h, txt, size=14, font=SANS, color=CREAM, bold=False, italic=False, align="left",
          spacing=None, leading=None, caps=False, strike=False):
        box = text(s, x, y, w, h, txt.upper() if caps else txt, font_size=size, color=color, bold=bold,
                   align=align, font_name=font, C=self.C)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for p in tf.paragraphs:
            p.alignment = tf.paragraphs[0].alignment
            if leading:
                p.line_spacing = leading
            for r in p.runs:
                r.font.italic = italic
                if spacing is not None:
                    r.font._element.set("spc", str(spacing))
                if strike:
                    r.font._element.set("strike", "sngStrike")
        return box

    def serif(self, s, x, y, w, h, txt, size=44, color=CREAM, italic=False, bold=True, align="left", leading=0.92):
        return self.T(s, x, y, w, h, txt, size=size, font=SERIF, color=color, bold=bold, italic=italic, align=align,
                      leading=leading)

    def shadow(self, shp, blur=24, dist=6, alpha=22):
        sp = shp._element.spPr
        eff = etree.SubElement(sp, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur * 12700), dist=str(dist * 12700),
                              dir="5400000", algn="t", rotWithShape="0")
        c = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(c, qn("a:alpha"), val=str(alpha * 1000))
        return shp

    def rbox(self, s, x, y, w, h, fill, radius=0.12, line=None, line_w=1.0, shadow=False):
        b = rrect(s, x, y, w, h, fill, line=line, C=self.C)
        b.adjustments[0] = radius
        if line:
            b.line.width = Pt(line_w)
        if shadow:
            self.shadow(b)
        return b

    def button(self, s, x, y, w, h, label, fill=RED, color=WHITE, size=14):
        b = self.rbox(s, x, y, w, h, fill, radius=0.5, shadow=True)
        self.T(s, x, y + (h - size / 72 * 1.3) / 2, w, h, label, size=size, color=color, bold=True, align="center")
        return b

    def sticker(self, s, x, y, d, label, fill=GOLD, color=INK, rot=-12, size=11):
        o = oval(s, x, y, d, d, fill, C=self.C)
        o.rotation = rot
        self.shadow(o, blur=14, dist=4, alpha=25)
        lines = label.split("\n")
        lh = size / 72 * 1.05
        top = y + d / 2 - lh * len(lines) / 2
        for k, ln in enumerate(lines):                       # one box per line: rotated multi-line boxes misalign
            t = self.T(s, x, top + k * lh, d, lh, ln, size=size, color=color, bold=True, align="center")
            t.rotation = rot
        return o

    def sparkle(self, s, cx, cy, d, fill):
        sp = shape(s, MSO_SHAPE.STAR_4_POINT, cx - d / 2, cy - d / 2, d, d, fill, C=self.C)
        sp.adjustments[0] = 0.18
        return sp

    def picture(self, s, name, x, y, w=None, h=None, tag="static-art"):
        pic = s.shapes.add_picture(str(ART / name), Inches(x), Inches(y), Inches(w) if w else None,
                                   Inches(h) if h else None)
        pic.name = tag
        return pic

    def slide(self, dark=True, notes=None):
        self.bg = BLACK if dark else CREAM
        s = add_slide(self.prs)
        self.picture(s, "grain_black.jpg" if dark else "grain_cream.jpg", 0, 0, w=W, h=H, tag="static-bg")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self):
        for i, sl in enumerate(self.prs.slides, start=1):
            motion.transition(sl, "fade", through_black=i in (1, 3))
            motion.choreograph(sl, step_ms=120, budget_ms=2200)
        OUT.mkdir(exist_ok=True)
        path = OUT / "Bella_by_Red_Flag_Homes.pptx"
        clean_save(self.prs, str(path))
        print(path)


def build():
    d = Deck()
    T, serif = d.T, d.serif

    # 1 — Cover (black) -----------------------------------------------------------------------------
    s = d.slide(notes="Lead with the problem Bella solves: the 3 AM guest message. Then the offer: 48 hours free.")
    d.picture(s, "logo.png", M - 0.1, 0.4, w=1.05)
    T(s, M, 1.75, 5, 0.3, "Red Flag Homes presents", size=11, color=RED, bold=True, caps=True, spacing=300)
    serif(s, M - 0.1, 1.9, 6.6, 3.2, "Bella", size=210, italic=True, leading=0.9)
    dot = oval(s, 5.95, 4.15, 0.36, 0.36, RED, C=d.C)
    d.sparkle(s, 5.85, 2.55, 0.42, GOLD)
    d.sparkle(s, 6.3, 3.0, 0.2, CREAM)
    rect(s, 6.75, 2.15, 0.012, 2.6, HAIR_D, C=d.C)
    serif(s, 7.0, 2.1, 2.8, 0.6, "What is Bella?", size=28, italic=True, color=CREAM)
    T(s, 7.0, 2.85, 2.3, 2.4, "An AI-powered agent by Red Flag Homes that runs your bnb for you — guest replies, "
      "pricing, check-ins and listing audits, 24/7.", size=12, color=CREAM, leading=1.45)
    T(s, 7.0, 4.75, 2.3, 0.8, "No per-booking commission.\nJust one yearly subscription.", size=11.5, color=GOLD,
      bold=True, leading=1.35)
    ax, ay, aw = 9.55, 1.35, 3.1                                            # arch portrait, offset red outline
    arch_h = aw * 1280 / 1220
    outline = shape(s, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, ax + 0.22, ay + 0.22, aw, arch_h, BLACK, line=RED, C=d.C)
    outline.adjustments[0] = 0.5
    outline.fill.background()
    outline.line.width = Pt(2)
    d.picture(s, "bella_arch.png", ax, ay, w=aw, h=arch_h)
    d.sticker(s, ax + aw - 0.75, ay + arch_h - 0.55, 1.2, "24/7", fill=RED, color=WHITE, size=19, rot=-10)
    d.sticker(s, ax + aw - 0.55, ay - 0.3, 1.0, "AI\nagent", fill=GOLD, color=INK, size=12, rot=12)
    d.sparkle(s, ax + aw * 0.45, ay + arch_h + 0.55, 0.32, CREAM)
    d.button(s, M, 5.75, 4.9, 0.66, "Start your 48-hour free trial  →", size=15)
    T(s, M + 5.2, 5.83, 6.5, 0.3, "Festive offer  ₹2,999 / year", size=14, color=CREAM, bold=True)
    T(s, M + 5.2, 6.15, 6.5, 0.3, "Exclusive to our first clients", size=10.5, color=MUTE_D)
    T(s, M, 6.95, 6, 0.25, "Bella by Red Flag Homes", size=8, color=MUTE_D, bold=True, caps=True, spacing=250)

    # 2 — What Bella does (cream) ------------------------------------------------------------------------
    s = d.slide(dark=False, notes="The chat on the phone is illustrative. Point at the 3:02 AM timestamp.")
    T(s, M, 0.7, 6, 0.3, "What Bella does", size=11, color=RED, bold=True, caps=True, spacing=300)
    serif(s, M, 0.95, 8.0, 1.6, "Everything your listing\nneeds, handled.", size=50, color=INK, leading=0.95)
    feats = [("01", "Instant guest replies", "Answers every guest in seconds, 24/7 — even at 3 AM."),
             ("02", "Smart pricing", "Adjusts your nightly rates for weekends, festivals and demand."),
             ("03", "Automatic check-in", "Sends every guest their check-in details on time, every time."),
             ("04", "A–Z listing audit", "Reviews photos, title, description and settings — and fixes the gaps.")]
    cw, ch, g = 3.55, 1.78, 0.22
    for i, (n, t, b) in enumerate(feats):
        x = M + (i % 2) * (cw + g)
        y = 2.75 + (i // 2) * (ch + g)
        dark = i == 0
        d.rbox(s, x, y, cw, ch, INK if dark else WHITE, radius=0.08, shadow=True)
        T(s, x + 0.3, y + 0.28, 0.8, 0.3, n, size=11, color=RED if not dark else GOLD, bold=True)
        serif(s, x + 0.3, y + 0.55, cw - 0.6, 0.5, t, size=24, color=CREAM if dark else INK)
        T(s, x + 0.3, y + 1.05, cw - 0.6, 0.65, b, size=10.5, color=BLUSH if dark else "#4D464A", leading=1.35)
    d.rbox(s, M, 6.75, 4.2, 0.42, RED, radius=0.5)
    T(s, M, 6.84, 4.2, 0.3, "Zero setup hassle.  Live in 3 clicks.", size=11.5, color=WHITE, bold=True, align="center")
    px, py, pw, ph = 9.0, 0.75, 3.35, 6.25                                   # phone with a 3 AM chat
    phone = d.rbox(s, px, py, pw, ph, INK, radius=0.13, shadow=True)
    d.rbox(s, px + 0.13, py + 0.13, pw - 0.26, ph - 0.26, "#FBF8F4", radius=0.11)
    d.rbox(s, px + pw / 2 - 0.5, py + 0.24, 1.0, 0.2, INK, radius=0.5)
    d.picture(s, "bella_avatar.png", px + 0.35, py + 0.6, w=0.55)
    T(s, px + 1.0, py + 0.64, 2, 0.25, "Bella", size=12, color=INK, bold=True)
    T(s, px + 1.0, py + 0.9, 2.2, 0.25, "● online · Red Flag Homes", size=8, color="#1DAA5B")
    rect(s, px + 0.13, py + 1.3, pw - 0.26, 0.012, HAIR_L, C=d.C)
    T(s, px, py + 1.45, pw, 0.25, "3:02 AM", size=8, color=MUTE, align="center", bold=True)
    g1 = d.rbox(s, px + 0.3, py + 1.8, 2.3, 0.75, "#ECE6DE", radius=0.25)
    T(s, px + 0.45, py + 1.92, 2.05, 0.6, "Hi! Can we check in early tomorrow?", size=9.5, color=INK, leading=1.3)
    b1 = d.rbox(s, px + 0.75, py + 2.75, 2.3, 1.1, RED, radius=0.2)
    T(s, px + 0.9, py + 2.87, 2.05, 0.95, "Hi! Early check-in from 11 AM is available. I’ve sent your check-in "
      "details 🔑", size=9.5, color=WHITE, leading=1.3)
    T(s, px + 0.75, py + 3.92, 2.3, 0.2, "Replied in 4 seconds", size=7.5, color=MUTE, align="right")
    g2 = d.rbox(s, px + 0.3, py + 4.25, 1.8, 0.5, "#ECE6DE", radius=0.3)
    T(s, px + 0.45, py + 4.37, 1.6, 0.3, "Amazing, thank you!", size=9.5, color=INK)
    d.sparkle(s, px - 0.25, py + 1.7, 0.4, RED)
    d.sparkle(s, px + pw + 0.12, py + 3.3, 0.26, GOLD)
    T(s, px, py + ph + 0.12, pw, 0.25, "Illustrative conversation", size=8, color=MUTE, align="center")

    # 3 — Pricing (black) -------------------------------------------------------------------------------
    s = d.slide(notes="Show the struck regular price first, then the festive card. Finish on ≈ ₹125 a month per listing.")
    d.picture(s, "logo.png", W - M - 0.9, 0.35, w=0.9)
    T(s, M, 0.75, 6, 0.3, "Pricing", size=11, color=RED, bold=True, caps=True, spacing=300)
    serif(s, M, 1.05, 10.2, 1.0, "No commission. Just one yearly fee.", size=42, italic=True)
    lx, ly = M, 2.65                                                        # regular price, struck
    d.rbox(s, lx, ly, 4.6, 3.1, CARD_D, radius=0.07, line=HAIR_D)
    T(s, lx + 0.4, ly + 0.4, 3, 0.3, "Regular price", size=11, color=MUTE_D, bold=True, caps=True, spacing=200)
    T(s, lx + 0.4, ly + 0.95, 4, 0.9, "₹4,999", size=54, color=MUTE_D, bold=True, strike=True)
    T(s, lx + 0.4, ly + 1.85, 3, 0.3, "per year", size=13, color=MUTE_D)
    T(s, lx + 0.4, ly + 2.4, 3.8, 0.3, "Covers up to 2 listings", size=12, color=MUTE_D)
    rx, ry, rw, rh = 5.75, 2.25, 4.9, 3.9                                   # festive offer, red card
    d.rbox(s, rx, ry, rw, rh, RED, radius=0.07, shadow=True)
    tag = d.rbox(s, rx + 0.4, ry + 0.4, 1.9, 0.38, BLACK, radius=0.5)
    T(s, rx + 0.4, ry + 0.47, 1.9, 0.3, "FESTIVE OFFER", size=9.5, color=GOLD, bold=True, align="center", spacing=200)
    T(s, rx + 0.4, ry + 1.0, 4.3, 1.2, "₹2,999", size=80, color=WHITE, bold=True)
    T(s, rx + 0.45, ry + 2.25, 3, 0.35, "per year", size=15, color=WHITE)
    rect(s, rx + 0.45, ry + 2.8, rw - 0.9, 0.012, "#E0707F", C=d.C)
    T(s, rx + 0.45, ry + 2.98, rw - 0.9, 0.3, "Covers up to 2 listings", size=13, color=WHITE, bold=True)
    T(s, rx + 0.45, ry + 3.32, rw - 0.9, 0.3, "48-hour free trial first", size=12, color=BLUSH)
    d.sticker(s, rx + rw - 0.85, ry - 0.55, 1.35, "Save\n₹2,000", fill=GOLD, color=INK, size=14, rot=14)
    d.picture(s, "bella_avatar.png", 11.05, 4.35, w=1.55)
    ring = oval(s, 11.0, 4.3, 1.65, 1.65, BLACK, line=GOLD, C=d.C)
    ring.fill.background()
    ring.line.width = Pt(1.5)
    d.sparkle(s, 12.7, 4.2, 0.3, GOLD)
    serif(s, M, 6.3, 10, 0.5, "That’s about ₹125 a month per listing with two listings — and a 48-hour free trial first.",
          size=19, italic=True, color=CREAM, bold=False)

    # 4 — How to start (cream) ---------------------------------------------------------------------------
    s = d.slide(dark=False, notes="Walk the three steps, then land on the button. Repeat: exclusive to our first clients.")
    T(s, M, 0.7, 6, 0.3, "How to start", size=11, color=RED, bold=True, caps=True, spacing=300)
    serif(s, M, 0.95, 11, 1.0, "Three steps. Then Bella takes over.", size=50, color=INK)
    steps = [("1", "Start your free trial", "48 hours free, up to 2 listings.", None),
             ("2", "Add Bella as co-host", "Invite her with full access:", "bella@redflaghomes.in"),
             ("3", "Save. Done.", "Bella now manages everything.", None)]
    sw, sg = 3.75, 0.29
    for i, (n, t, b, chip) in enumerate(steps):
        x = M + i * (sw + sg)
        y = 2.35
        d.rbox(s, x, y, sw, 2.75, WHITE if i < 2 else INK, radius=0.07, shadow=True)
        T(s, x + 0.35, y + 0.25, 1.2, 1.2, n, size=72, color=RED, bold=True)
        serif(s, x + 0.35, y + 1.4, sw - 0.7, 0.5, t, size=26, color=INK if i < 2 else CREAM)
        T(s, x + 0.35, y + 1.95, sw - 0.7, 0.4, b, size=11.5, color="#4D464A" if i < 2 else BLUSH)
        if chip:
            d.rbox(s, x + 0.35, y + 2.25, 2.75, 0.34, "#F7E3E6", radius=0.5)
            T(s, x + 0.35, y + 2.31, 2.75, 0.25, chip, size=10, color=RED, bold=True, align="center")
        if i < 2:
            T(s, x + sw - 0.05, y + 1.1, 0.4, 0.4, "→", size=20, color=RED, bold=True, align="center")
    d.sparkle(s, 12.45, 2.15, 0.36, RED)
    d.sparkle(s, 12.75, 2.55, 0.2, GOLD)
    strip = rect(s, 0, 5.6, W, 1.9, BLACK, C=d.C)
    d.picture(s, "bella_avatar.png", M, 5.85, w=1.3)
    serif(s, M + 1.6, 5.98, 6, 0.6, "Exclusive to Red Flag’s first clients.", size=26, italic=True, color=CREAM)
    T(s, M + 1.6, 6.55, 6, 0.3, "Festive offer ₹2,999 / year  ·  48-hour free trial", size=11, color=MUTE_D)
    d.button(s, W - M - 3.3, 6.0, 3.3, 0.8, "Get Bella  →", size=18)
    d.picture(s, "logo.png", W - M - 4.45, 5.85, w=0.95)

    assert len(d.prs.slides) == 4
    d.save()


if __name__ == "__main__":
    build()
